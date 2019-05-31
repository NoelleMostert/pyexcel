# import openpyxl

# wb = openpyxl.load_workbook('example.xlsx')
# sheet = wb.get_sheet_by_name('Sheet1')
# tuple(sheet['A1':'C3'])
# for rowOfCellObjects in sheet['A1':'C3']:
#     for cellObj in rowOfCellObjects:
#         print(cellObj.value)
#     print('--- END OF ROW ---')

import openpyxl
wb = openpyxl.load_workbook('example.xlsx')
sheet = wb.active
for i in range(1, 3):
        sheet['A' + str(i)] = i

refObj = openpyxl.chart.Reference(sheet, min_col=1, min_row=1, max_col=12, max_row=3)

seriesObj = openpyxl.chart.Series(refObj, title='First series')

chartObj = openpyxl.chart.BarChart()
chartObj.title = 'My Chart'
chartObj.append(seriesObj)
sheet.add_chart(chartObj, 'C5')
wb.save('sampleChart.xlsx')